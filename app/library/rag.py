from __future__ import annotations

import os
import json
import uuid
import shutil
# pyrefly: ignore [missing-import]
import chromadb
from pathlib import Path
import hashlib
import redis
from rank_bm25 import BM25Okapi
from ..core.config import settings

# Paths
BASE_DIR = Path(__file__).resolve().parent.parent.parent
KNOWLEDGE_BASE_DIR = BASE_DIR / "knowledge_base"
DATA_DIR = BASE_DIR / "data"
DATA_DIR.mkdir(parents=True, exist_ok=True)
KNOWLEDGE_BASE_DIR.mkdir(parents=True, exist_ok=True)

# A persistent folder for files uploaded via the Admin Panel
PERSISTENT_KB_DIR = DATA_DIR / "kb_uploads"
PERSISTENT_KB_DIR.mkdir(parents=True, exist_ok=True)

CHROMA_DB_DIR = DATA_DIR / "chroma_db"

# Supported file extensions
SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".md", ".csv", ".json", ".docx", ".xlsx"}

# Initialize ChromaDB persistent client
chroma_client = chromadb.PersistentClient(path=str(CHROMA_DB_DIR))

# Get or create the collection
collection = chroma_client.get_or_create_collection(name="library_knowledge_base")

# ── BM25 Hybrid Search State ─────────────────────────────────────────────────
_bm25_index = None
_bm25_corpus = []

def rebuild_bm25_index():
    """Reads all current chunks from ChromaDB and rebuilds the BM25 index in memory."""
    global _bm25_index, _bm25_corpus
    all_data = collection.get()
    
    if all_data and all_data.get("documents") and len(all_data["documents"]) > 0:
        _bm25_corpus = [
            {
                "id": all_data["ids"][i],
                "text": all_data["documents"][i],
                "metadata": all_data["metadatas"][i]
            }
            for i in range(len(all_data["documents"]))
        ]
        # Tokenize by splitting on whitespace for BM25
        tokenized = [doc["text"].lower().split() for doc in _bm25_corpus]
        _bm25_index = BM25Okapi(tokenized)
        print(f"[RAG] Rebuilt BM25 index with {len(_bm25_corpus)} documents.")
    else:
        _bm25_index = None
        _bm25_corpus = []
        print("[RAG] Cleared BM25 index (no documents).")

# Initialize BM25 on startup from existing DB chunks
rebuild_bm25_index()

# Initialize Redis client lazily so a Redis outage never crashes the app
_redis_client = None

def get_redis() -> redis.Redis | None:
    """Return a Redis client, or None if Redis is not reachable."""
    global _redis_client
    if _redis_client is not None:
        return _redis_client
    try:
        client = redis.Redis.from_url(settings.REDIS_URL, decode_responses=True, socket_connect_timeout=2)
        client.ping()  # Verify the connection is actually alive
        _redis_client = client
        print(f"[Redis] Connected successfully to {settings.REDIS_URL}")
        return _redis_client
    except Exception as e:
        print(f"[Redis] Could not connect ({e}). Caching disabled — falling back to ChromaDB.")
        return None




# ── Text extractors ──────────────────────────────────────────────────────────

def _extract_txt(file_path: Path) -> str:
    """Extract text from a plain text or markdown file."""
    try:
        return file_path.read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        print(f"[RAG] Error reading {file_path.name}: {e}")
        return ""


def _extract_pdf(file_path: Path) -> str:
    """Extract text from a PDF file using pypdf."""
    try:
        # pyrefly: ignore [missing-import]
        from pypdf import PdfReader
        reader = PdfReader(str(file_path))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        return "\n\n".join(pages)
    except ImportError:
        print("[RAG] pypdf not installed. Cannot process PDF files.")
        return ""
    except Exception as e:
        print(f"[RAG] Error reading PDF {file_path.name}: {e}")
        return ""


def _extract_docx(file_path: Path) -> str:
    """Extract text from a .docx file using python-docx."""
    try:
        # pyrefly: ignore [missing-import]
        from docx import Document
        doc = Document(str(file_path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)
    except ImportError:
        print("[RAG] python-docx not installed. Cannot process .docx files.")
        return ""
    except Exception as e:
        print(f"[RAG] Error reading DOCX {file_path.name}: {e}")
        return ""


def _extract_csv(file_path: Path) -> str:
    """Extract text from a CSV file."""
    try:
        import csv
        rows = []
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            for row in reader:
                rows.append(", ".join(row))
        return "\n".join(rows)
    except Exception as e:
        print(f"[RAG] Error reading CSV {file_path.name}: {e}")
        return ""


def _extract_json(file_path: Path) -> str:
    """Extract text from a JSON file by serializing it as formatted text."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            data = json.load(f)
        return json.dumps(data, indent=2, ensure_ascii=False)
    except Exception as e:
        print(f"[RAG] Error reading JSON {file_path.name}: {e}")
        return ""


def _extract_xlsx(file_path: Path) -> str:
    """Extract text from an Excel .xlsx file using openpyxl."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        all_rows = []
        for sheet in wb.worksheets:
            all_rows.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                row_text = ", ".join(str(v) for v in row if v is not None)
                if row_text.strip():
                    all_rows.append(row_text)
        return "\n".join(all_rows)
    except ImportError:
        print("[RAG] openpyxl not installed. Cannot process .xlsx files.")
        return ""
    except Exception as e:
        print(f"[RAG] Error reading XLSX {file_path.name}: {e}")
        return ""


def extract_text(file_path: Path) -> str:
    """Dispatch to the correct extractor based on file extension."""
    ext = file_path.suffix.lower()
    if ext in (".txt", ".md"):
        return _extract_txt(file_path)
    elif ext == ".pdf":
        return _extract_pdf(file_path)
    elif ext == ".docx":
        return _extract_docx(file_path)
    elif ext == ".csv":
        return _extract_csv(file_path)
    elif ext == ".json":
        return _extract_json(file_path)
    elif ext == ".xlsx":
        return _extract_xlsx(file_path)
    else:
        print(f"[RAG] Unsupported file type: {ext}")
        return ""


# ── Chunking ─────────────────────────────────────────────────────────────────

def sliding_window_chunker(text: str, chunk_size: int = 1000, chunk_overlap: int = 200) -> list[str]:
    """Chunks text using a sliding window based on character count."""
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunks.append(text[start:end])
        if end == len(text):
            break
        start += (chunk_size - chunk_overlap)
    return chunks


# ── Core ingestion ────────────────────────────────────────────────────────────

def ingest_documents(chunk_size: int = 1000, chunk_overlap: int = 200):
    """
    Reads all supported files in the knowledge_base directory and upserts them into ChromaDB.
    Supports: .txt, .md, .pdf, .docx, .csv, .json, .xlsx
    """
    if not KNOWLEDGE_BASE_DIR.exists():
        print(f"Directory not found: {KNOWLEDGE_BASE_DIR}")
        return

    documents = []
    metadatas = []
    ids = []

    for file_path in KNOWLEDGE_BASE_DIR.iterdir():
        if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        content = extract_text(file_path)
        if not content.strip():
            print(f"[RAG] No content extracted from {file_path.name}, skipping.")
            continue

        chunks = sliding_window_chunker(content, chunk_size=chunk_size, chunk_overlap=chunk_overlap)
        for i, chunk in enumerate(chunks):
            if not chunk.strip():
                continue
            doc_id = f"{file_path.stem}_{i}"
            documents.append(chunk)
            metadatas.append({"source": file_path.name, "chunk_index": i, "file_type": file_path.suffix.lower()})
            ids.append(doc_id)

    if documents:
        collection.upsert(documents=documents, metadatas=metadatas, ids=ids)
        print(f"[RAG] Ingested {len(documents)} chunks from {KNOWLEDGE_BASE_DIR}.")
        rebuild_bm25_index()
    else:
        print("[RAG] No documents found to ingest.")
        rebuild_bm25_index()


def add_document_to_kb(filename: str, content: str, chunk_size: int = 1000, chunk_overlap: int = 200):
    """Chunks a document's content and upserts it into ChromaDB."""
    chunks = sliding_window_chunker(content, chunk_size, chunk_overlap)
    if not chunks:
        return

    # First remove old chunks for this file to allow clean re-upload
    _delete_file_chunks(filename)

    stem = Path(filename).stem
    ext = Path(filename).suffix.lower()
    documents = []
    metadatas = []
    ids = []

    for i, chunk in enumerate(chunks):
        if not chunk.strip():
            continue
        doc_id = f"{stem}_{i}"
        documents.append(chunk)
        metadatas.append({"source": filename, "chunk_index": i, "file_type": ext})
        ids.append(doc_id)

    if documents:
        collection.upsert(documents=documents, metadatas=metadatas, ids=ids)
        print(f"[RAG] Added {len(documents)} chunks for '{filename}'.")
        rebuild_bm25_index()


def _delete_file_chunks(filename: str):
    """Remove all ChromaDB chunks that belong to a specific file."""
    try:
        existing = collection.get(where={"source": filename})
        if existing and existing["ids"]:
            collection.delete(ids=existing["ids"])
            print(f"[RAG] Removed {len(existing['ids'])} old chunks for '{filename}'.")
            rebuild_bm25_index()
    except Exception as e:
        print(f"[RAG] Warning: Could not remove old chunks for '{filename}': {e}")


def delete_document(filename: str) -> bool:
    """Remove a document file from disk and its chunks from ChromaDB."""
    file_path = KNOWLEDGE_BASE_DIR / filename
    if not file_path.exists():
        return False
    file_path.unlink()
    _delete_file_chunks(filename)
    return True


def list_documents() -> list[dict]:
    """Return a list of all documents in the knowledge_base directories with metadata."""
    docs = []
    
    # Read from GitHub synced directory
    for file_path in sorted(KNOWLEDGE_BASE_DIR.iterdir()):
        if file_path.suffix.lower() in SUPPORTED_EXTENSIONS:
            docs.append({
                "filename": file_path.name,
                "size_bytes": file_path.stat().st_size,
                "file_type": file_path.suffix.lower().lstrip(".").upper(),
                "source": "github"
            })
            
    # Read from Persistent Admin Uploads directory
    for file_path in sorted(PERSISTENT_KB_DIR.iterdir()):
        if file_path.suffix.lower() in SUPPORTED_EXTENSIONS:
            docs.append({
                "filename": file_path.name,
                "size_bytes": file_path.stat().st_size,
                "file_type": file_path.suffix.lower().lstrip(".").upper(),
                "source": "admin_upload"
            })
            
    return docs


def clear_knowledge_base():
    """Deletes all documents from the knowledge base collection."""
    existing = collection.get()
    if existing and existing["ids"]:
        collection.delete(ids=existing["ids"])
        rebuild_bm25_index()


def search_knowledge_base(query: str, n_results: int = 3) -> str:
    """
    Search the library knowledge base for the given query.
    Returns a JSON string containing the most relevant text passages.
    """
    try:
        # Normalize the query
        normalized_query = query.strip().lower()
        
        # Generate a cache key
        query_hash = hashlib.sha256(normalized_query.encode('utf-8')).hexdigest()
        cache_key = f"rag_cache:{query_hash}"
        
        # Check Redis cache (get_redis() returns None if Redis is unavailable)
        rc = get_redis()
        if rc:
            try:
                cached_result = rc.get(cache_key)
                if cached_result:
                    print(f"[Cache Hit] Returning cached result for query: '{query}'")
                    return cached_result
            except Exception as redis_err:
                print(f"[Redis Error] Failed to read from cache: {redis_err}")
            
        print(f"[Cache Miss] Querying Hybrid Search for: '{query}'")
        
        # 1. Vector Search (ChromaDB)
        vector_results = collection.query(query_texts=[query], n_results=n_results * 2)
        vector_scores = {}
        if vector_results and vector_results["ids"]:
            for i, doc_id in enumerate(vector_results["ids"][0]):
                distance = vector_results["distances"][0][i]
                if distance < 1.5:  # Semantic distance filter
                    vector_scores[doc_id] = 1.0 / (60 + i + 1)

        # 2. Keyword Search (BM25)
        bm25_scores = {}
        if _bm25_index is not None and len(_bm25_corpus) > 0:
            tokenized_query = query.lower().split()
            bm25_raw_scores = _bm25_index.get_scores(tokenized_query)
            # Get indices of top scores
            top_indices = sorted(range(len(bm25_raw_scores)), key=lambda idx: bm25_raw_scores[idx], reverse=True)[:n_results * 2]
            for rank, idx in enumerate(top_indices):
                if bm25_raw_scores[idx] > 0:
                    doc_id = _bm25_corpus[idx]["id"]
                    bm25_scores[doc_id] = 1.0 / (60 + rank + 1)

        # 3. Reciprocal Rank Fusion (RRF)
        rrf_scores = {}
        all_ids = set(vector_scores.keys()).union(set(bm25_scores.keys()))
        for doc_id in all_ids:
            rrf_scores[doc_id] = vector_scores.get(doc_id, 0.0) + bm25_scores.get(doc_id, 0.0)

        # Sort by RRF score and take top n_results
        sorted_ids = sorted(rrf_scores.keys(), key=lambda k: rrf_scores[k], reverse=True)[:n_results]

        passages = []
        if sorted_ids:
            final_docs = collection.get(ids=sorted_ids)
            # Map retrieved docs back to their sorted order
            id_to_doc = {
                final_docs["ids"][i]: {
                    "source": final_docs["metadatas"][i].get("source", "Unknown"),
                    "text": final_docs["documents"][i]
                }
                for i in range(len(final_docs["ids"]))
            }
            for doc_id in sorted_ids:
                if doc_id in id_to_doc:
                    passages.append(id_to_doc[doc_id])

        if passages:
            response_json = json.dumps({"status": "success", "results": passages})
        else:
            response_json = json.dumps({"status": "not_found", "message": "No relevant information found in the knowledge base."})
            
        # Store in Redis with TTL (e.g., 3600 seconds = 1 hour)
        if rc:
            try:
                rc.setex(cache_key, 3600, response_json)
            except Exception as redis_err:
                print(f"[Redis Error] Failed to write to cache: {redis_err}")
            
        return response_json

    except Exception as e:
        return json.dumps({"status": "error", "message": str(e)})


if __name__ == "__main__":
    ingest_documents()
